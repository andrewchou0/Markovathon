"""Generate Markov.pptx — the deck as a real PowerPoint file.

Fonts are Arial Narrow (display) and Georgia (body): both ship with Office on
Windows and macOS, so the file will not substitute fonts on an unknown laptop.
Google-hosted faces would look better and render as something else on the day.

python-pptx has no animation API, so slide transitions are injected as raw XML
and element builds are left to PowerPoint's own Animations tab.
"""
from pptx import Presentation
from pptx.util import Inches as In, Pt, Emu
from pptx.dml.color import RGBColor as C
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
import copy

W, H = In(13.333), In(7.5)
GROUND, SURFACE, SURFACE2 = C(0x10, 0x13, 0x0F), C(0x17, 0x1B, 0x15), C(0x1E, 0x23, 0x19)
INK, INK2, MUTED, LINE = C(0xE8, 0xEA, 0xE3), C(0xC3, 0xC8, 0xBA), C(0x8D, 0x94, 0x86), C(0x2B, 0x31, 0x26)
BRASS, BRASS_DIM, BRASS_LIFT = C(0xC8, 0xA0, 0x4A), C(0x8E, 0x71, 0x33), C(0xE3, 0xC4, 0x83)
OXIDE, OXIDE_DIM = C(0xD4, 0x55, 0x3F), C(0x7D, 0x2F, 0x22)
TIERS = [C(0xE3, 0xC4, 0x83), C(0xC8, 0xA0, 0x4A), C(0x9C, 0x7A, 0x33), C(0x6D, 0x55, 0x22)]
DISPLAY, BODY, MONO = "Arial Narrow", "Georgia", "Consolas"

prs = Presentation()
prs.slide_width, prs.slide_height = W, H
BLANK = prs.slide_layouts[6]
MARGIN = In(0.72)
CONTENT_W = W - 2 * MARGIN


def slide():
    s = prs.slides.add_slide(BLANK)
    bg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, W, H)
    bg.fill.solid(); bg.fill.fore_color.rgb = GROUND; bg.line.fill.background()
    bg.shadow.inherit = False
    # marking bar: defense documents carry one, and ours declares the data fictional
    bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, W, In(0.26))
    bar.fill.solid(); bar.fill.fore_color.rgb = OXIDE_DIM; bar.line.fill.background()
    bar.shadow.inherit = False
    write(bar.text_frame, "UNCLASSIFIED  //  ILLUSTRATIVE FIGURES  //  CONTAINS NO CONTROLLED TECHNICAL DATA",
          MONO, 9, C(0xF4, 0xE6, 0xE2), align=PP_ALIGN.CENTER, space=1.6)
    bar.text_frame.margin_top = bar.text_frame.margin_bottom = 0
    return s


def write(tf, text, font, size, color, bold=False, align=PP_ALIGN.LEFT, space=0, caps=False, line=None):
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    if line:
        p.line_spacing = line
    r = p.add_run(); r.text = text.upper() if caps else text
    r.font.name, r.font.size, r.font.bold = font, Pt(size), bold
    r.font.color.rgb = color
    if space:
        rPr = r._r.get_or_add_rPr(); rPr.set("spc", str(int(space * 100)))
    return p


def box(s, x, y, w, h, fill=None, line_color=None, line_w=1.0, top_accent=None, left_accent=None):
    sh = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
    if fill is None:
        sh.fill.background()
    else:
        sh.fill.solid(); sh.fill.fore_color.rgb = fill
    if line_color is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line_color; sh.line.width = Pt(line_w)
    sh.shadow.inherit = False
    if top_accent:
        a = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, Pt(2.6))
        a.fill.solid(); a.fill.fore_color.rgb = top_accent; a.line.fill.background(); a.shadow.inherit = False
    if left_accent:
        a = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, Pt(3), h)
        a.fill.solid(); a.fill.fore_color.rgb = left_accent; a.line.fill.background(); a.shadow.inherit = False
    return sh


def text(s, x, y, w, h, content, font, size, color, **kw):
    tb = s.shapes.add_textbox(x, y, w, h)
    tb.text_frame.margin_left = tb.text_frame.margin_right = 0
    tb.text_frame.margin_top = tb.text_frame.margin_bottom = 0
    write(tb.text_frame, content, font, size, color, **kw)
    return tb


def eyebrow(s, idx, label, total=11):
    text(s, MARGIN, In(0.62), CONTENT_W, In(0.3),
         f"{idx:02d}   {label.upper()}   / {total}", MONO, 10.5, BRASS, space=1.8)


def heading(s, content, y=In(1.02), size=34, h=In(1.5), color=INK):
    return text(s, MARGIN, y, CONTENT_W, h, content, DISPLAY, size, color, bold=True, caps=True, line=0.95)


def cite(s, content):
    text(s, MARGIN, H - In(0.78), CONTENT_W, In(0.45), content.upper(), MONO, 8.5, MUTED, space=1.2, line=1.35)


def paras(s, x, y, w, items, size=13, color=INK2, gap=In(0.28), font=BODY, line=1.35):
    """Stacked paragraphs in one textbox, so spacing can't double up."""
    tb = s.shapes.add_textbox(x, y, w, In(0.4))
    tf = tb.text_frame; tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, (txt, kw) in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.line_spacing = kw.get("line", line)
        p.space_after = Pt(kw.get("after", 8))
        r = p.add_run(); r.text = txt.upper() if kw.get("caps") else txt
        r.font.name = kw.get("font", font); r.font.size = Pt(kw.get("size", size))
        r.font.bold = kw.get("bold", False); r.font.color.rgb = kw.get("color", color)
        if kw.get("space"):
            r._r.get_or_add_rPr().set("spc", str(int(kw["space"] * 100)))
    return tb


# ============================== 01 ==============================
s = slide()
# wordmark: four ascending bars — a chart, and a chain of states
bx = MARGIN
for i, (hh, col) in enumerate([(In(0.22), BRASS_DIM), (In(0.34), BRASS), (In(0.46), BRASS_LIFT), (In(0.58), OXIDE)]):
    b = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, bx + In(0.115) * i, In(1.02) - hh, In(0.075), hh)
    b.fill.solid(); b.fill.fore_color.rgb = col; b.line.fill.background(); b.shadow.inherit = False
text(s, MARGIN + In(0.56), In(0.52), In(3), In(0.55), "MARKOV", DISPLAY, 32, INK, bold=True, space=1)
tag = box(s, MARGIN + In(2.35), In(0.6), In(3.5), In(0.32), line_color=BRASS_DIM)
write(tag.text_frame, "BUILT IN ONE DAY  ·  GB10  ·  NO CLOUD API", MONO, 9, BRASS, align=PP_ALIGN.CENTER, space=1.4)
eyebrow(s, 1, "Why we built this")
heading(s, "Everyone shipped a cloud agent.", y=In(1.5), size=40, h=In(0.72))
heading(s, "The people who need one most are not allowed to use it.", y=In(2.16), size=40, h=In(1.5), color=BRASS)
paras(s, MARGIN, In(3.9), In(7.4), [
    ("This year every useful AI agent became a request to somebody else's datacenter. That is fine if you sell shoes. The American defense industrial base cannot send a supplier list, a part number, or a program schedule to a commercial model — not as policy, as law.", {"after": 12}),
    ("So the organizations carrying the most fragile supply chains in the country were handed exactly none of this decade's tooling. We built the agent they are permitted to run: it never leaves the machine it is installed on.", {}),
], size=14)
cite(s, "Governing instruments: DFARS 252.204-7012 · NIST SP 800-171 · CMMC 2.0 · ITAR 22 CFR 120-130")

# ============================== 02 ==============================
s = slide()
eyebrow(s, 2, "Naming it properly")
heading(s, 'We do not build for "disruptions." A disruption is weather.', size=33, h=In(1.2))
d = box(s, MARGIN, In(2.32), CONTENT_W, In(1.5), fill=SURFACE, left_accent=BRASS)
text(s, MARGIN + In(0.34), In(2.56), CONTENT_W - In(0.7), In(0.4), "Plan-invalidating event", DISPLAY, 23, BRASS_LIFT, bold=True, caps=True)
text(s, MARGIN + In(0.34), In(3.04), CONTENT_W - In(0.7), In(0.7),
     "Any event that makes the plan you are currently executing impossible to execute as written — and therefore obliges someone to decide, quickly, what replaces it.",
     BODY, 15, INK, line=1.35)
paras(s, MARGIN, In(4.08), In(8.6), [
    ("The distinction is not stylistic. “Disruption” describes a condition in the world. A plan-invalidating event describes a state change in your obligations: a schedule you can no longer meet, a delivery you can no longer promise, a clause you are now in breach of.", {}),
], size=14)
colw = (CONTENT_W - In(0.6)) / 3
for i, (h3, bodytxt) in enumerate([
    ("The event", "An export licence is suspended, indefinitely, on a rare-earth concentrate."),
    ("What it invalidates", "Not a shipment. The build sequence for every assembly drawing that specifies that magnet billet."),
    ("Who has to decide", "A named human, inside hours, with authority to commit — who must first be told which programs are exposed."),
]):
    x = MARGIN + (colw + In(0.3)) * i
    text(s, x, In(5.15), colw, In(0.3), h3, DISPLAY, 12.5, BRASS, bold=True, caps=True, space=1.6)
    text(s, x, In(5.5), colw, In(1.0), bodytxt, BODY, 12.5, INK2, line=1.35)
cite(s, "None of the three examples we demonstrate are weather. All three stop a plan.")

# ============================== 03 ==============================
s = slide()
eyebrow(s, 3, "The market")
heading(s, "A hundred thousand companies, and the one that stops the program is invisible", size=31, h=In(1.2))
tiles = [("100,000+", "Companies and subcontractors under contract in the defense industrial base.", "CISA, DIB sector"),
         ("~300,000", "Companies and suppliers in the broader defense ecosystem.", "DoD CIO"),
         ("12,000+", "Small and mid-size subcontractors whose capacities — sometimes whose identities — are not well known to DoD.", "Brookings"),
         ("51 → 5", "Aerospace and defense prime contractors, after three decades of consolidation.", "DoD, 2022")]
tw = (CONTENT_W - In(0.36)) / 4
for i, (fig, what, src) in enumerate(tiles):
    x = MARGIN + (tw + In(0.12)) * i
    box(s, x, In(2.3), tw, In(2.3), fill=SURFACE)
    text(s, x + In(0.2), In(2.48), tw - In(0.4), In(0.6), fig, DISPLAY, 32, BRASS_LIFT, bold=True)
    text(s, x + In(0.2), In(3.08), tw - In(0.4), In(1.1), what, BODY, 11.5, INK2, line=1.3)
    text(s, x + In(0.2), In(4.24), tw - In(0.4), In(0.3), src.upper(), MONO, 8, MUTED, space=1.2)
kb = box(s, MARGIN, In(4.86), CONTENT_W, In(1.15), left_accent=BRASS)
text(s, MARGIN + In(0.26), In(4.92), CONTENT_W - In(0.5), In(1.0),
     "Read the third tile again. By the Department's own account, thousands of the firms that programs depend on are not reliably identified. You cannot manage exposure through a supplier you cannot name — and the vendor base shrank from roughly 69,000 firms to 55,000 between 2016 and 2020, so each remaining one carries more.",
     BODY, 13.5, INK, line=1.4)
cite(s, "CISA DIB sector profile · DoD CIO · Brookings · DoD State of Competition within the Defense Industrial Base, Feb 2022")

# ============================== 04 ==============================
s = slide()
eyebrow(s, 4, "The annual bill")
heading(s, "What a year of unexecutable plans actually costs", size=34, h=In(0.8))
tiles = [("$1.4T", "Lost annually to unplanned downtime across the world's 500 largest companies — 11% of total revenue, up from 8% in 2019.", "Siemens, 2024"),
         ("$82M", "Average annual losses per company from supply-side events.", "Interos, 2022"),
         ("$184M", "Average lost revenue per year, per large organization.", "Interos"),
         ("$22M", "Average cost of a single event, across surveyed companies.", "Interos")]
for i, (fig, what, src) in enumerate(tiles):
    x = MARGIN + (tw + In(0.12)) * i
    box(s, x, In(1.92), tw, In(2.2), fill=SURFACE)
    text(s, x + In(0.2), In(2.1), tw - In(0.4), In(0.6), fig, DISPLAY, 34, OXIDE, bold=True)
    text(s, x + In(0.2), In(2.72), tw - In(0.4), In(1.1), what, BODY, 11.5, INK2, line=1.3)
    text(s, x + In(0.2), In(3.76), tw - In(0.4), In(0.3), src.upper(), MONO, 8, MUTED, space=1.2)
for i, (h3, bodytxt) in enumerate([
    ("Revenue", "Deliveries missed inside the quarter they were promised in. On a defense program, that is a milestone payment, not a deferred sale."),
    ("Labor", "A stopped line does not stop payroll. Cleared, trained staff stand idle at full cost — then the recovery is paid again at overtime."),
    ("Trust", "Past performance is scored. A subcontractor that surprises a prime twice stops being asked to bid."),
]):
    x = MARGIN + (colw + In(0.3)) * i
    text(s, x, In(4.42), colw, In(0.3), h3, DISPLAY, 12.5, BRASS, bold=True, caps=True, space=1.6)
    text(s, x, In(4.78), colw, In(1.2), bodytxt, BODY, 12.5, INK2, line=1.35)
cite(s, "Separate studies measuring different populations, deliberately not plotted on one scale. Siemens True Cost of Downtime 2024 · Interos supplier-risk research")

# ============================== 05 ==============================
s = slide()
eyebrow(s, 5, "The minute")
heading(s, "A large plant loses $33,000 a minute. The clock starts before anyone knows why.", size=33, h=In(1.2))
box(s, MARGIN, In(2.5), CONTENT_W, In(2.3), fill=SURFACE, top_accent=OXIDE)
text(s, MARGIN + In(0.4), In(2.74), CONTENT_W - In(0.8), In(0.3), "ACCRUING AT $33,000 / MINUTE", MONO, 10, MUTED, space=1.8)
text(s, MARGIN + In(0.4), In(3.0), CONTENT_W - In(0.8), In(1.1), "$1,650 every 3 seconds", DISPLAY, 58, OXIDE, bold=True)
text(s, MARGIN + In(0.4), In(4.12), In(8.6), In(0.5),
     "One stopped station backs up welding, paint, final assembly, labor and outbound shipping at once. A full plant hour runs to $2.3 million.",
     BODY, 13, INK2, line=1.35)
kb = box(s, MARGIN, In(5.05), CONTENT_W, In(1.0), left_accent=BRASS)
text(s, MARGIN + In(0.26), In(5.1), CONTENT_W - In(0.5), In(0.9),
     "Average time to restore production has risen from 49 minutes to 81. Very little of that is the physical fix. It is the interval spent establishing which suppliers are affected, how far down the chain it reaches, and who has the authority to act. That interval is information work — and it is the part we remove.",
     BODY, 13.5, INK, line=1.4)
cite(s, "Siemens, True Cost of Downtime 2024. Per-minute figure as reported; the hourly figure implies ~$38,000, so $33,000 is the conservative of the two.")

# ============================== 06 ==============================
s = slide()
eyebrow(s, 6, "Why it must be local")
heading(s, "Sending this to a cloud model is not a privacy preference. It may be an unlicensed export.", size=31, h=In(1.3))
for i, (h3, bodytxt) in enumerate([
    ("Person-based, not geographic", "ITAR restricts access by foreign persons regardless of where the server sits. If a provider's infrastructure is operated by or accessible to foreign nationals, sending controlled technical data to it can constitute an export without a licence."),
    ("“US region” is not a control", "A provider subject to foreign government access law creates sovereignty exposure whatever the datacenter's address. Routing CUI through a commercial tool's external servers violates the contract directly."),
    ("The penalty is personal", "Up to $1M per violation, with criminal liability reaching executives. No procurement officer signs off on that to save an afternoon of analysis."),
]):
    x = MARGIN + (colw + In(0.3)) * i
    text(s, x, In(2.42), colw, In(0.34), h3, DISPLAY, 13, BRASS, bold=True, caps=True, space=1.6)
    text(s, x, In(2.82), colw, In(1.7), bodytxt, BODY, 12.5, INK2, line=1.4)
box(s, MARGIN, In(4.72), CONTENT_W, In(1.5), fill=SURFACE, left_accent=BRASS)
text(s, MARGIN + In(0.34), In(4.94), CONTENT_W - In(0.7), In(0.34), "And the graph is the secret", DISPLAY, 20, BRASS_LIFT, bold=True, caps=True)
text(s, MARGIN + In(0.34), In(5.36), CONTENT_W - In(0.7), In(0.8),
     "A supplier network is not metadata about a program. Which firm machines which part, sole-sourced, for which assembly — that IS the program, described from the bottom up. An adversary who wants to know what you are building, and where to press, would ask for exactly this file.",
     BODY, 14, INK, line=1.35)
cite(s, "ITAR 22 CFR 120-130 · DFARS 252.204-7012 · NIST SP 800-171 · CMMC 2.0 · CUI enclave guidance")

# ============================== 07 ==============================
s = slide()
eyebrow(s, 7, "What we built, and how")
heading(s, "It watches unattended, traces three tiers, and drafts the response. The model writes — it never decides who is affected.", size=27, h=In(1.25))
stages = [("01 · local", "MongoDB", "Suppliers, parts, sole-source flags, the dependency graph. Loopback only."),
          ("02 · deterministic", "propagation.py", "Who is hit and how far it cascades. Plain Python, cycle-safe, same answer every run."),
          ("03 · openclaw", "Gateway → Qwen", "The gateway runs the prompt. The model writes; it is never asked which suppliers matter."),
          ("04 · openclaw", "Channel → reviewer", "The draft is delivered to the person who can approve it, in the app they already use."),
          ("05 · human", "Approval", "Unsent until a person says so. Every decision, including suppressions, is logged.")]
sw = (CONTENT_W - In(0.48)) / 5
for i, (n, name, desc) in enumerate(stages):
    x = MARGIN + (sw + In(0.12)) * i
    accent = BRASS_LIFT if i == 1 else BRASS_DIM
    box(s, x, In(2.42), sw, In(1.72), fill=SURFACE, top_accent=accent)
    text(s, x + In(0.16), In(2.6), sw - In(0.32), In(0.2), n.upper(), MONO, 7.5, MUTED, space=1.2)
    text(s, x + In(0.16), In(2.82), sw - In(0.32), In(0.26), name, MONO, 11, INK, bold=True)
    text(s, x + In(0.16), In(3.14), sw - In(0.32), In(0.9), desc, BODY, 10, INK2, line=1.3)
    if i < 4:
        ar = s.shapes.add_shape(MSO_SHAPE.RIGHT_TRIANGLE, x + sw + Emu(6000), In(3.16), In(0.11), In(0.14))
        ar.rotation = 90; ar.fill.solid(); ar.fill.fore_color.rgb = BRASS_DIM
        ar.line.fill.background(); ar.shadow.inherit = False
# cascade ladder — depth is magnitude, one hue stepped light to dark
rungs = [("tier 0", "Altiplano Rare Earth", "neodymium magnet billet · sole source · non-compliant", 1.0),
         ("tier 1", "2 suppliers", "sensor array · mounting bracket", 0.84),
         ("tier 2", "2 suppliers", "turbocharger subassembly · wiring harness", 0.64),
         ("tier 3", "Finished powertrain module", "the deliverable", 0.44)]
y = In(4.34)
for i, (tag, name, note, frac) in enumerate(rungs):
    t_ = box(s, MARGIN, y, In(0.86), In(0.4), fill=SURFACE2)
    write(t_.text_frame, tag.upper(), MONO, 9, INK2, align=PP_ALIGN.CENTER, space=1.2)
    barw = Emu(int((CONTENT_W - In(0.96)) * frac))
    b = box(s, MARGIN + In(0.96), y, barw, In(0.4), fill=TIERS[i])
    ink = GROUND if i < 2 else C(0xF2, 0xEE, 0xE4)
    tb = s.shapes.add_textbox(MARGIN + In(1.1), y + In(0.03), barw - In(0.3), In(0.34))
    tf = tb.text_frame; tf.word_wrap = False
    tf.margin_left = tf.margin_top = tf.margin_bottom = 0
    p = tf.paragraphs[0]
    r1 = p.add_run(); r1.text = name.upper() + "   "
    r1.font.name, r1.font.size, r1.font.bold, r1.font.color.rgb = DISPLAY, Pt(12), True, ink
    r2 = p.add_run(); r2.text = note
    r2.font.name, r2.font.size, r2.font.color.rgb = BODY, Pt(9.5), ink
    y += In(0.46)
cite(s, "Worked example from the running system: one licence suspension, six firms exposed, network risk 0.961 — four tiers from the product that ships.   Measured: 0.57s with the model off · 267 automated checks · external calls made: 0")

# ============================== 08 ==============================
s = slide()
eyebrow(s, 8, "Built with")
heading(s, "Eleven components. Every one of them on the box.", size=34, h=In(0.8))
cols = [In(0.5), In(2.3), In(6.6), In(1.6)]
hx = MARGIN
for label, cw in zip(["REF", "COMPONENT", "FUNCTION", "RUNS AT"], cols):
    text(s, hx, In(1.92), cw, In(0.22), label, MONO, 8, BRASS_DIM, space=1.4,
         align=PP_ALIGN.RIGHT if label == "RUNS AT" else PP_ALIGN.LEFT)
    hx += cw + In(0.06)
rows = [("A1", "NVIDIA GB10", "The box. Everything below this line runs on it.", "on prem", False),
        ("A2", "Qwen 3.6 35B", "Language model. Writes the assessment and the draft.", "local weights", False),
        ("A3", "Ollama", "Serves the model on the box.", ":11434", False),
        ("A4", "OpenClaw", "Agent gateway. Two jobs — next slide.", ":18789", True),
        ("A5", "MongoDB 7", "The supplier graph and the event log.", ":27017", False),
        ("A6", "FastAPI + uvicorn", "Five endpoints, thin routing, no logic.", ":8000", False),
        ("A7", "React 19 + Vite", "The operations workspace.", ":5173", False),
        ("B1", "propagation.py", "Ours. Deterministic exposure. No model, no network.", "in process", True),
        ("B2", "monitor.py", "Ours. The unattended loop, thresholds, audit trail.", "daemon thread", True),
        ("B3", "narrate.py", "Ours. Sanitises, gates, falls back so text always arrives.", "in process", True),
        ("B4", "267 checks", "143 agent core · 77 end-to-end · 30 handoff · 17 frontend.", "green", True)]
y = In(2.2)
rh = In(0.355)
for ref, item, role, where, ours in rows:
    x = MARGIN
    fill = SURFACE2 if ours else SURFACE
    for j, (val, cw) in enumerate(zip([ref, item, role, where], cols)):
        box(s, x, y, cw, rh - Emu(18000), fill=fill)
        if j == 0:
            text(s, x, y + In(0.09), cw, In(0.2), val, MONO, 8, MUTED, align=PP_ALIGN.CENTER)
        elif j == 1:
            text(s, x + In(0.12), y + In(0.075), cw - In(0.2), In(0.22), val, MONO, 10,
                 BRASS_LIFT if ours else INK)
        elif j == 2:
            text(s, x + In(0.12), y + In(0.08), cw - In(0.2), In(0.22), val, BODY, 10.5, INK2)
        else:
            text(s, x + In(0.04), y + In(0.085), cw - In(0.14), In(0.2), val, MONO, 9, BRASS, align=PP_ALIGN.RIGHT)
        x += cw + In(0.06)
    y += rh
text(s, MARGIN, y + In(0.12), In(9.5), In(0.4),
     "Outbound internet connections required at run time: zero. The rows marked in brass are ours; the rest is off-the-shelf and stays on the machine.",
     BODY, 12.5, INK2, line=1.35)
cite(s, "Ports are loopback bindings. A non-local host for the model, the gateway or the database is refused before a socket opens.")

# ============================== 09 ==============================
s = slide()
eyebrow(s, 9, "OpenClaw, twice")
heading(s, "One component does the two jobs a local agent actually needs", size=34, h=In(0.85))
jobs = [("JOB 01 — MODEL HARNESS", "The model becomes a plug, not a hard-wire",
         ["We do not call the model directly. We post the prompt to the gateway's OpenAI-compatible endpoint and its agent resolves which model answers.",
          "POST /v1/chat/completions",
          "Its config names the model: ollama/qwen3.6:35b. Swapping to a larger model, or a different one entirely, is a line of config — not a code change."]),
        ("JOB 02 — APPROVAL CHANNEL", "The draft reaches a human where they already are",
         ["An agent that drafts a response for approval has to reach the person who approves. OpenClaw speaks Slack, Telegram, iMessage, Teams and more, from the box.",
          "POST /tools/invoke   x-openclaw-message-channel: slack",
          "No inbox to check, no dashboard to remember. The exposure summary and the draft arrive in the thread they already work in."])]
jw = (CONTENT_W - In(0.16)) / 2
for i, (n, title, lines) in enumerate(jobs):
    x = MARGIN + (jw + In(0.16)) * i
    box(s, x, In(2.12), jw, In(2.56), fill=SURFACE, left_accent=BRASS)
    text(s, x + In(0.28), In(2.3), jw - In(0.5), In(0.22), n, MONO, 8.5, BRASS_DIM, space=1.5)
    text(s, x + In(0.28), In(2.56), jw - In(0.5), In(0.5), title, DISPLAY, 16, INK, bold=True, caps=True, line=1.05)
    text(s, x + In(0.28), In(3.08), jw - In(0.5), In(0.6), lines[0], BODY, 11.5, INK2, line=1.35)
    cb = box(s, x + In(0.28), In(3.78), jw - In(0.56), In(0.3), fill=GROUND)
    text(s, x + In(0.36), In(3.84), jw - In(0.7), In(0.22), lines[1], MONO, 9, BRASS)
    text(s, x + In(0.28), In(4.16), jw - In(0.5), In(0.6), lines[2], BODY, 11.5, INK2, line=1.35)
box(s, MARGIN, In(4.92), CONTENT_W, In(1.1), left_accent=BRASS)
text(s, MARGIN + In(0.26), In(4.98), CONTENT_W - In(0.5), In(1.0),
     "And it is layered, never depended on. The chain is OpenClaw → direct Ollama → deterministic template: if the gateway is off, unauthorised, or has that endpoint disabled, the assessment still arrives. Adding a gateway must not add a way to fail.",
     BODY, 13.5, INK, line=1.4)
cite(s, "Gateway on 127.0.0.1:18789 · auth via bearer token · a non-local gateway host is refused without a socket being opened")

# ============================== 10 ==============================
s = slide()
eyebrow(s, 10, "What it saves")
heading(s, "We do not shorten the repair. We shorten the distance between the event and the decision.", size=31, h=In(1.2))
box(s, MARGIN, In(2.36), CONTENT_W, In(2.12), fill=SURFACE)
calc = [("Rate while the plan is unexecutable", "$33,000 / min", False),
        ("Restoration window today, average", "81 min", False),
        ("Decision latency we remove — assumed, ~12% of the window", "10 min", False),
        ("Avoided, per plan-invalidating event", "$330,000", True)]
cy = In(2.56)
for label, val, total in calc:
    if total:
        rule = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, MARGIN + In(0.3), cy - In(0.06), CONTENT_W - In(0.6), Pt(2))
        rule.fill.solid(); rule.fill.fore_color.rgb = BRASS; rule.line.fill.background(); rule.shadow.inherit = False
        cy += In(0.1)
    text(s, MARGIN + In(0.3), cy, In(7.4), In(0.3), label, BODY, 12.5, INK if total else INK2)
    text(s, W - MARGIN - In(2.9), cy - (In(0.07) if total else 0), In(2.6), In(0.34), val,
         DISPLAY if total else MONO, 20 if total else 12, BRASS_LIFT if total else INK,
         bold=total, align=PP_ALIGN.RIGHT)
    cy += In(0.46)
text(s, MARGIN + In(0.3), In(4.02), In(9.2), In(0.4),
     "Every input is on this slide so you can argue with the assumption rather than the total. Change the 10 minutes to 3 and it is $99,000. A single mid-size subcontractor seeing twelve such events a year, at ten minutes each, is $4.0M.",
     BODY, 11, MUTED, line=1.35)
rf = box(s, MARGIN, In(4.66), CONTENT_W, In(1.36), fill=C(0x19, 0x12, 0x0F), line_color=OXIDE_DIM)
text(s, MARGIN + In(0.26), In(4.8), CONTENT_W - In(0.5), In(0.24), "What we will not put on this slide", DISPLAY, 13, OXIDE, bold=True, caps=True, space=1.2)
text(s, MARGIN + In(0.26), In(5.1), CONTENT_W - In(0.5), In(0.8),
     "A national savings headline. The only defensible per-minute figure we can source is for a large plant; multiplying it across 12,000 small subcontractors would be arithmetic theatre, and you would be right to dismiss the rest of the deck for it. The macro anchor that IS sourced: 11% of revenue, up from 8% in five years. The national return is also schedule — and schedule slip on a fielded capability is measured in readiness, not dollars.",
     BODY, 11.5, INK2, line=1.35)
cite(s, "Rate and window: Siemens 2024. Per-event totals are modeled from those two inputs and labeled as such.")

# ============================== 11 ==============================
s = slide()
eyebrow(s, 11, "Beyond defense")
heading(s, "Three conditions. Wherever all three hold, this is the same product.", size=34, h=In(0.85))
for i, (n, h3, bodytxt) in enumerate([
    ("01", "A dependency graph", "Things depend on things, more than one layer deep, and nobody can hold the whole graph in their head."),
    ("02", "Events that invalidate plans", "Something happens that makes the standing plan unexecutable, and a human must choose the replacement quickly."),
    ("03", "Data that cannot leave", "Law, contract, or plain commercial sense forbids sending the graph to a third party's model."),
]):
    x = MARGIN + (colw + In(0.3)) * i
    text(s, x, In(2.2), colw, In(0.3), f"{n} · {h3}".upper(), DISPLAY, 13, BRASS, bold=True, space=1.4)
    text(s, x, In(2.58), colw, In(1.1), bodytxt, BODY, 12.5, INK2, line=1.4)
chips = ["Pharmaceutical · GxP, validated supply, sole-source APIs", "Semiconductor · tool chains, allocation, export control",
         "Energy · grid assets, long-lead transformers", "Hospital systems · PHI, single-source consumables",
         "Banking · vendor concentration under regulation"]
cx, cy2 = MARGIN, In(3.9)
for ch in chips:
    cwid = In(0.1) + In(0.062) * len(ch)
    if cx + cwid > W - MARGIN:
        cx, cy2 = MARGIN, cy2 + In(0.42)
    cb = box(s, cx, cy2, cwid, In(0.34), fill=SURFACE2)
    text(s, cx + In(0.12), cy2 + In(0.08), cwid - In(0.2), In(0.2), ch, MONO, 9, INK2)
    cx += cwid + In(0.1)
text(s, MARGIN, In(4.86), In(9.4), In(0.5),
     "Defense is where the constraint is sharpest, which is why we built it there first. The architecture does not soften anywhere else — it simply stops being mandatory and starts being an advantage.",
     BODY, 14, INK2, line=1.4)
kb = box(s, MARGIN, In(5.6), CONTENT_W, In(0.62), left_accent=BRASS)
text(s, MARGIN + In(0.26), In(5.72), CONTENT_W - In(0.5), In(0.4),
     "Ask: give us the stage, and we will run it live — insert an event into the database, touch nothing, and let the agent find it while you watch.",
     BODY, 14, INK, line=1.3)
cite(s, "Markov · built in one day, on one box, with no cloud API call at any point.")

# ---- slide transitions: python-pptx has no API, so inject the XML directly ----
# A single fade per slide. Element builds are left to PowerPoint's Animations tab;
# fabricating <p:timing> trees risks a file that will not open, which is a worse
# outcome than a deck the presenter animates in two clicks.
for sl in prs.slides:
    el = sl._element
    trans = el.makeelement(qn("p:transition"), {"spd": "med"})
    trans.append(trans.makeelement(qn("p:fade"), {}))
    el.append(trans)

prs.save("Markov.pptx")
print("wrote Markov.pptx")
print("slides:", len(prs.slides._sldIdLst))
